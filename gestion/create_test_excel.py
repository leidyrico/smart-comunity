import openpyxl
from openpyxl import Workbook
from datetime import datetime

# Crear un nuevo libro de Excel
wb = Workbook()

# Eliminar la hoja por defecto
wb.remove(wb.active)

# Crear hoja de Apartamentos
apartamentos_sheet = wb.create_sheet(title='Apartamentos')
apartamentos_headers = ['numero', 'piso', 'torre', 'propietario', 'telefono', 'email', 'area_m2', 'tipo', 'estado', 'estatus_financiero', 'observaciones']
apartamentos_sheet.append(apartamentos_headers)

# Datos de ejemplo para apartamentos
apartamentos_data = [
    ['701', '7', 'A', 'Carlos Rodriguez', '3001111111', 'carlos@email.com', '70.5', 'apartamento', 'ocupado', 'solvente', ''],
    ['702', '7', 'A', 'Ana Martinez', '3002222222', 'ana@email.com', '65.0', 'apartamento', 'ocupado', 'deudor', ''],
    ['703', '7', 'A', 'Luis Gomez', '3003333333', 'luis@email.com', '80.0', 'apartamento', 'ocupado', 'solvente', '']
]

for row in apartamentos_data:
    apartamentos_sheet.append(row)

# Crear hoja de Recibos
recibos_sheet = wb.create_sheet(title='Recibos')
recibos_headers = ['numero_recibo', 'periodo', 'fecha_emision', 'fecha_vencimiento', 'valor_administracion', 'valor_aseo', 'valor_vigilancia', 'valor_mantenimiento', 'otros_conceptos', 'estado']
recibos_sheet.append(recibos_headers)

# Datos de ejemplo para recibos
recibos_data = [
    ['REC-701-2024-01', '2024-01', '01-01-2024', '31-01-2024', '150000', '20000', '30000', '15000', '0', 'activo'],
    ['REC-701-2024-02', '2024-02', '01-02-2024', '29-02-2024', '150000', '20000', '30000', '15000', '5000', 'activo'],
    ['REC-702-2024-01', '2024-01', '01-01-2024', '31-01-2024', '140000', '18000', '28000', '12000', '0', 'activo'],
    ['REC-702-2024-02', '2024-02', '01-02-2024', '29-02-2024', '140000', '18000', '28000', '12000', '0', 'activo'],
    ['REC-703-2024-01', '2024-01', '01-01-2024', '31-01-2024', '160000', '22000', '32000', '18000', '0', 'activo'],
    ['REC-703-2024-02', '2024-02', '01-02-2024', '29-02-2024', '160000', '22000', '32000', '18000', '8000', 'activo']
]

for row in recibos_data:
    recibos_sheet.append(row)

# Crear hoja de Pagos
pagos_sheet = wb.create_sheet(title='Pagos')
pagos_headers = ['apartamento_numero', 'recibo_numero', 'fecha_pago', 'monto_pagado', 'metodo_pago', 'numero_comprobante', 'estado', 'observaciones']
pagos_sheet.append(pagos_headers)

# Datos de ejemplo para pagos
pagos_data = [
    ['701', 'REC-701-2024-01', '15-01-2024', '215000', 'transferencia', 'TRF-001', 'confirmado', 'Pago completo'],
    ['702', 'REC-702-2024-01', '20-01-2024', '100000', 'efectivo', 'EFE-001', 'confirmado', 'Pago parcial'],
    ['703', 'REC-703-2024-01', '25-01-2024', '232000', 'cheque', 'CHE-001', 'confirmado', 'Pago completo']
]

for row in pagos_data:
    pagos_sheet.append(row)

# Guardar el archivo
wb.save('test_import_completo.xlsx')
print('Archivo Excel creado exitosamente: test_import_completo.xlsx')